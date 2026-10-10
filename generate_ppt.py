import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette - Enterprise Modern Dark Theme
    COLOR_BG = RGBColor(11, 19, 41)         # #0b1329 deep navy
    COLOR_CARD = RGBColor(23, 37, 65)       # #172541 slate blue
    COLOR_CARD_BORDER = RGBColor(51, 65, 85) # #334155
    COLOR_PRIMARY = RGBColor(59, 130, 246)  # #3b82f6 royal blue
    COLOR_ACCENT = RGBColor(6, 182, 212)    # #06b6d4 cyan
    COLOR_SUCCESS = RGBColor(16, 185, 129)  # #10b981 emerald
    COLOR_WARNING = RGBColor(245, 158, 11)  # #f59e0b amber
    COLOR_DANGER = RGBColor(239, 68, 68)    # #ef4444 crimson
    COLOR_TEXT_MAIN = RGBColor(248, 250, 252) # #f8fafc white
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8 slate gray

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="FRAUDSHIELD AI ENTERPRISE PLATFORM"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN

    def add_card(slide, left, top, width, height, title, items, border_color=COLOR_PRIMARY, title_color=COLOR_TEXT_MAIN):
        # Background container shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tx_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(17)
        p_title.font.bold = True
        p_title.font.color.rgb = title_color
        p_title.space_after = Pt(12)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.space_after = Pt(8)

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative glow card
    center_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(10.9), Inches(5.1))
    center_card.fill.solid()
    center_card.fill.fore_color.rgb = COLOR_CARD
    center_card.line.color.rgb = COLOR_PRIMARY
    center_card.line.width = Pt(2)

    tb1 = s1.shapes.add_textbox(Inches(1.8), Inches(1.7), Inches(9.7), Inches(4.1))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "ENTERPRISE SOFTWARE ENGINEERING | ADVANCED JAVA"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_ACCENT
    p_badge.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "AI-Powered Fraud Detection System"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_TEXT_MAIN
    p_main.space_after = Pt(12)

    p_sub = tf1.add_paragraph()
    p_sub.text = "High-Performance Autonomous Payment Security & Real-Time Risk Analytics"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED
    p_sub.space_after = Pt(28)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Pure Standard Java | Polymorphic Engine Architecture | Asynchronous Pipeline | Pure JDBC Pooling"
    p_desc.font.size = Pt(13)
    p_desc.font.color.rgb = COLOR_PRIMARY

    # ==========================================
    # SLIDE 2: THE PROBLEM & MOTIVATION
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Financial Challenge: Fraud at Scale")

    add_card(s2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "The $38B Annual Threat",
             [
                 "Global online payment fraud costs financial institutions over $38B annually.",
                 "Sophisticated attacks: automated card velocity testing, identity spoofing, and wire laundering.",
                 "Legacy batch processing fails to catch fast fraud before settlement occurs."
             ], border_color=COLOR_DANGER, title_color=COLOR_DANGER)

    add_card(s2, Inches(4.85), Inches(1.6), Inches(3.6), Inches(5.2),
             "The False Positive Dilemma",
             [
                 "Rigid threshold rules block legitimate customers, causing 15%+ checkout abandonment.",
                 "Customer friction leads to direct lost revenue and brand damage.",
                 "Need for explainable multivariate scoring over binary blunt blocking."
             ], border_color=COLOR_WARNING, title_color=COLOR_WARNING)

    add_card(s2, Inches(8.9), Inches(1.6), Inches(3.6), Inches(5.2),
             "Our Architectural Solution",
             [
                 "Sub-10ms decision latency with dual-strategy hybrid evaluation.",
                 "Explainable AI scoring providing transparency to compliance teams.",
                 "Zero framework overhead: Pure Java concurrency and thread-safe collections."
             ], border_color=COLOR_SUCCESS, title_color=COLOR_SUCCESS)

    # ==========================================
    # SLIDE 3: SYSTEM ARCHITECTURE OVERVIEW
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Enterprise 4-Tier Architectural Blueprint")

    tiers = [
        ("1. Presentation Tier", Inches(0.8), [
            "Jakarta EE Servlets & dynamic JSP views",
            "Responsive Bootstrap 5 dark-mode console",
            "Role-based routing: Admin vs User portals",
            "Real-time status badges & triage controls"
        ], COLOR_ACCENT),
        ("2. Concurrency Pipeline", Inches(3.85), [
            "Producer-Consumer bounded LinkedBlockingQueue",
            "Managed ThreadPoolExecutor daemon workers",
            "Lock-free metrics: AtomicInteger & DoubleAdder",
            "In-memory sliding velocity cache (ConcurrentHashMap)"
        ], COLOR_PRIMARY),
        ("3. AI Detection Engines", Inches(6.9), [
            "Factory Pattern supplying dynamic engines",
            "RuleBasedEngine (velocity, geo, limit heuristics)",
            "MLAnomalyEngine (sigmoid isolation scoring)",
            "HybridEngine ensemble combining both models"
        ], COLOR_WARNING),
        ("4. Pure JDBC Persistence", Inches(9.95), [
            "Custom DBConnectionPool (thread-safe, health checks)",
            "Exclusive PreparedStatement usage (zero SQL injection)",
            "Manual ACID transactions: setAutoCommit(false)",
            "GenericDAO<T, ID> full CRUD contract"
        ], COLOR_SUCCESS)
    ]

    for title, left, bullets, color in tiers:
        add_card(s3, left, Inches(1.6), Inches(2.7), Inches(5.2), title, bullets, border_color=color, title_color=color)

    # ==========================================
    # SLIDE 4: OOP & DESIGN PATTERNS
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Object-Oriented Design & Applied Patterns")

    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.45),
             "Inheritance & Polymorphism",
             [
                 "Abstract base class 'Transaction' with polymorphic contracts.",
                 "'CreditCardTransaction': Card network, CVV match, billing zip.",
                 "'WireTransaction': SWIFT/BIC formatting, beneficiary jurisdiction.",
                 "Polymorphic baseline risk calculation tailored per payment rail."
             ], border_color=COLOR_PRIMARY)

    add_card(s4, Inches(6.8), Inches(1.6), Inches(5.6), Inches(2.45),
             "Factory Pattern (EngineFactory)",
             [
                 "Decouples client controllers from concrete detection algorithms.",
                 "Dynamic runtime switching: 'RULE_BASED', 'ML_ANOMALY', 'HYBRID'.",
                 "Dynamic configuration changes take effect without server restart."
             ], border_color=COLOR_ACCENT)

    add_card(s4, Inches(0.8), Inches(4.35), Inches(5.6), Inches(2.45),
             "Custom Exception Hierarchy",
             [
                 "Root checked exception: 'FraudDetectionException'.",
                 "'DatabaseException': Encapsulates SQL errors & pool timeouts.",
                 "'ThresholdOutOfBoundsException': Enforces strict risk boundary constraints."
             ], border_color=COLOR_WARNING)

    add_card(s4, Inches(6.8), Inches(4.35), Inches(5.6), Inches(2.45),
             "Generic DAO Pattern (GenericDAO<T, ID>)",
             [
                 "Generic CRUD interface decouples persistence from services.",
                 "Implemented by UserDAO, TransactionDAO, SystemConfigDAO, FraudAlertDAO.",
                 "Zero ORM bloat: 100% pure JDBC type-safe mapping."
             ], border_color=COLOR_SUCCESS)

    # ==========================================
    # SLIDE 5: DETECTION ENGINES COMPARISON
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Dual-Model Engine: Deterministic vs AI Heuristic")

    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Deterministic Rule-Based Engine",
             [
                 "Velocity Anomalies: Detects high transaction frequency within sliding time window (e.g., >3 tx in 10 mins).",
                 "Geographic Checks: Flags high-risk OFAC/FATF foreign destinations or unresolved IP locations.",
                 "High-Amount Deviation: Checks variance against baseline spending limits ($5,000 threshold).",
                 "Technical Rail Checks: CVV mismatch penalty (+0.30) and malformed SWIFT code penalty (+0.25).",
                 "Deterministic Output: Explicit triggered rule audit trails."
             ], border_color=COLOR_PRIMARY, title_color=COLOR_PRIMARY)

    add_card(s5, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Machine Learning Anomaly Engine",
             [
                 "Multivariate Feature Normalization [0.0 - 1.0]:",
                 "  • f_amount: Sigmoid non-linear logistic amount curve",
                 "  • f_velocity: Normalized transaction frequency density",
                 "  • f_geo: Entropy score for cross-border destinations",
                 "  • f_device: Unrecognized hardware fingerprint entropy",
                 "  • f_rail: Intrinsic base payment rail risk factor",
                 "Dynamic Weight Aggregation: Sum(w_i * f_i) / Sum(w_i)",
                 "Explainable AI: Outputs transparent feature contribution weights."
             ], border_color=COLOR_ACCENT, title_color=COLOR_ACCENT)

    # ==========================================
    # SLIDE 6: MULTITHREADING & ASYNC PIPELINE
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "High-Throughput Multithreading & Concurrency")

    add_card(s6, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Producer-Consumer Pipeline",
             [
                 "Bounded LinkedBlockingQueue<Transaction> handles traffic spikes.",
                 "Thread backpressure prevents out-of-memory crashes.",
                 "Dedicated ThreadPoolExecutor daemon workers process queue in background.",
                 "Sub-millisecond enqueuing allows instant UI response."
             ], border_color=COLOR_PRIMARY)

    add_card(s6, Inches(4.85), Inches(1.6), Inches(3.6), Inches(5.2),
             "Thread-Safe Collections",
             [
                 "ConcurrentHashMap: Lock-free reads for system configuration & velocity maps.",
                 "CopyOnWriteArrayList: Thread-safe sliding transaction history per user.",
                 "PriorityBlockingQueue<FraudAlert>: Automatically surfaces CRITICAL blocked threats ahead of minor warnings."
             ], border_color=COLOR_ACCENT)

    add_card(s6, Inches(8.9), Inches(1.6), Inches(3.6), Inches(5.2),
             "Lock-Free Metrics Tracking",
             [
                 "AtomicInteger: Contention-free tracking for processed, approved, and blocked counts.",
                 "DoubleAdder: High-concurrency financial volume and risk accumulation.",
                 "Real-time accuracy & false-positive precision tracking in live UI."
             ], border_color=COLOR_SUCCESS)

    # ==========================================
    # SLIDE 7: PURE JDBC & POOLING
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Pure JDBC Persistence & Manual ACID Control")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Custom DBConnectionPool",
             [
                 "Zero third-party pooling frameworks: Built with pure Java java.sql.*.",
                 "Thread-Safe Queue: BlockingQueue<Connection> manages pool lifecycle.",
                 "Active Health Check: Verifies connection.isValid(2) before distribution.",
                 "Leak Prevention: Auto-recovers stale or uncommitted connections.",
                 "Dynamic Schema Init: Auto-executes schema.sql and seeds initial data on startup."
             ], border_color=COLOR_PRIMARY, title_color=COLOR_PRIMARY)

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Strict Manual ACID Transactions",
             [
                 "connection.setAutoCommit(false): Explicit transactional boundaries.",
                 "Atomic Multi-Table Inserts: Persists Transaction record and FraudAlert in same unit of work.",
                 "Explicit Commit: connection.commit() commits state only upon complete success.",
                 "Safe Rollback: connection.rollback() in catch blocks guarantees zero orphaned records.",
                 "100% PreparedStatement: Eliminates SQL injection vulnerabilities."
             ], border_color=COLOR_SUCCESS, title_color=COLOR_SUCCESS)

    # ==========================================
    # SLIDE 8: ADMIN DASHBOARD & CONTROLS
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Admin Oversight & Dynamic Risk Tuning")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Dynamic Risk Boundaries",
             [
                 "Flag Threshold (Review): Defaults to 0.70 (escalates to human triage).",
                 "Block Threshold (Hard Stop): Defaults to 0.90 (halts payment immediately).",
                 "Strict Validation: Rejects invalid bounds (< 0.0 or Flag >= Block)."
             ], border_color=COLOR_PRIMARY)

    add_card(s8, Inches(4.85), Inches(1.6), Inches(3.6), Inches(5.2),
             "Real-Time ML Weight Tuning",
             [
                 "Live sliders for Amount, Location, Device, and Velocity weights.",
                 "Hot-swaps active detection engine between Rule, ML, and Hybrid.",
                 "Immediate persistence to system_configs table and sync across threads."
             ], border_color=COLOR_ACCENT)

    add_card(s8, Inches(8.9), Inches(1.6), Inches(3.6), Inches(5.2),
             "Human-in-the-Loop Triage",
             [
                 "Filterable Live Ledger: Inspect Approved, Under Review, or Blocked.",
                 "Analyst Actions: One-click Approve, Block, or Investigate.",
                 "False-Positive Feedback Loop: Approving a flagged transaction increments accuracy calibration counters."
             ], border_color=COLOR_WARNING)

    # ==========================================
    # SLIDE 9: USER EXPERIENCE & WORKFLOW
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "User Experience & Payment Workflow")

    add_card(s9, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "1. Payment Initiation",
             [
                 "Tabbed Payment Rails: Credit Card vs High-Value Wire Transfer.",
                 "Form captures amount, currency, recipient/card, location, device fingerprint, and category.",
                 "Option to test instant synchronous evaluation or async queue worker pipeline."
             ], border_color=COLOR_PRIMARY)

    add_card(s9, Inches(4.85), Inches(1.6), Inches(3.6), Inches(5.2),
             "2. Instant Decision Feedback",
             [
                 "Color-Coded Status Badges: Green (APPROVED), Amber (UNDER_REVIEW), Red (BLOCKED).",
                 "Risk Score Gauge: Shows exact confidence percentage.",
                 "Explainable AI Diagnostic: Displays top anomaly features responsible for the decision."
             ], border_color=COLOR_SUCCESS)

    add_card(s9, Inches(8.9), Inches(1.6), Inches(3.6), Inches(5.2),
             "3. Security Incident Notices",
             [
                 "Dedicated Alerts Inbox prioritized by severity (CRITICAL, HIGH, MEDIUM, LOW).",
                 "Alerts linked directly to suspicious transactions.",
                 "Audit timestamps and user acknowledgment status."
             ], border_color=COLOR_DANGER)

    # ==========================================
    # SLIDE 10: AUTOMATED TESTING & VERIFICATION
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Automated Verification & Test Suite")

    tests = [
        ("Test Suite 1: OOP & Polymorphism", "Validates base risk factors for CreditCard vs Wire subclasses.", COLOR_PRIMARY),
        ("Test Suite 2: Detection Engines", "Validates Rule-Based and ML Anomaly scoring against safe & high-risk profiles.", COLOR_ACCENT),
        ("Test Suite 3: Generic DAO CRUD", "Verifies pure JDBC persistence, user querying, and schema consistency.", COLOR_SUCCESS),
        ("Test Suite 4: Custom Exceptions", "Enforces ThresholdOutOfBoundsException validation on invalid risk inputs.", COLOR_WARNING),
        ("Test Suite 5: Async Pipeline", "Verifies background worker consumption, persistence, and queue thread-safety.", COLOR_PRIMARY),
        ("Test Suite 6: Priority Alert Triage", "Verifies PriorityBlockingQueue triage ordering (CRITICAL > MEDIUM > LOW).", COLOR_DANGER),
        ("Test Suite 7: False-Positive Calibration", "Verifies metrics synchronization when an analyst overrides a flagged status.", COLOR_ACCENT)
    ]

    for idx, (t_title, t_desc, t_col) in enumerate(tests):
        y_pos = Inches(1.6 + (idx * 0.72))
        bar = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.7), Inches(0.62))
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_CARD
        bar.line.color.rgb = t_col
        bar.line.width = Pt(1)

        tb = s10.shapes.add_textbox(Inches(1.0), y_pos + Inches(0.08), Inches(11.3), Inches(0.48))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        run1 = p.add_run()
        run1.text = "✔ " + t_title + " — "
        run1.font.bold = True
        run1.font.size = Pt(13)
        run1.font.color.rgb = t_col

        run2 = p.add_run()
        run2.text = t_desc
        run2.font.size = Pt(12)
        run2.font.color.rgb = COLOR_TEXT_MAIN

    # ==========================================
    # SLIDE 11: DEPLOYMENT & PRODUCTION READINESS
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Deployment & Cloud Infrastructure")

    add_card(s11, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Multi-Stage Docker Container",
             [
                 "Stage 1: Maven + JDK 21 build container packages JAR & lib/.",
                 "Stage 2: Lightweight Alpine JRE runtime image.",
                 "Exposes dynamic $PORT environment variable for cloud host compatibility."
             ], border_color=COLOR_PRIMARY)

    add_card(s11, Inches(4.85), Inches(1.6), Inches(3.6), Inches(5.2),
             "Instant Public Bridge",
             [
                 "TLS-terminated public HTTPS tunnel bridge.",
                 "Accessible globally without opening local firewall ports.",
                 "Cross-device mobile-responsive Bootstrap 5 interface."
             ], border_color=COLOR_ACCENT)

    add_card(s11, Inches(8.9), Inches(1.6), Inches(3.6), Inches(5.2),
             "1-Click Cloud Hosting Ready",
             [
                 "render.yaml: Automated infrastructure-as-code for Render web services.",
                 "Procfile: Ready for Railway, Fly.io, and Heroku deployment.",
                 "Git repository initialized with clean commit history."
             ], border_color=COLOR_SUCCESS)

    # ==========================================
    # SLIDE 12: CONCLUSION & SUMMARY
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Summary & Enterprise Engineering Takeaways")

    add_card(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Core Engineering Competencies Mastered",
             [
                 "Robust OOP Abstraction: Polymorphic models and custom exception hierarchy.",
                 "Industrial Multithreading: High-throughput Producer-Consumer queue with lock-free atomic telemetry.",
                 "Pure Enterprise JDBC: Custom connection pool with manual ACID transaction rollback/commit integrity.",
                 "Explainable AI: Multivariate continuous anomaly scoring with dynamic administrative tuning."
             ], border_color=COLOR_PRIMARY, title_color=COLOR_PRIMARY)

    add_card(s12, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "System Highlights & Credentials",
             [
                 "Zero Heavy Framework Overhead: 100% standard Java and Jakarta Servlets.",
                 "Instant 1-Click Launch: run.bat / mvnw compile exec:java.",
                 "Admin Access: admin / admin123 (Full control & metrics).",
                 "User Access: john_doe / password123 (Transactions & alerts).",
                 "Live Public Deployment: Bridged via secure HTTPS."
             ], border_color=COLOR_SUCCESS, title_color=COLOR_SUCCESS)

    output_file = "Fraud_Detection_System_Presentation.pptx"
    prs.save(output_file)
    print(f"Presentation saved successfully to: {output_file}")

if __name__ == "__main__":
    create_deck()
